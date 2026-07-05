from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUTPUT_PATH = Path("FastAPI_Backend_Documentation.docx")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_code(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(10)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True)
        shade_cell(table.rows[0].cells[index], "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
    doc.add_paragraph()


def build_document() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)
    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        doc.styles[style_name].font.name = "Times New Roman"
        doc.styles[style_name].font.bold = True

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("FastAPI Backend Documentation")
    run.bold = True
    run.font.size = Pt(20)
    run.font.name = "Times New Roman"

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Battery Remaining Useful Life Prediction System")
    run.font.size = Pt(14)
    run.font.name = "Times New Roman"

    doc.add_paragraph()

    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(
        "This document explains the FastAPI backend developed for the Battery Remaining Useful Life "
        "(RUL) prediction system. The FastAPI backend acts as the connection layer between the "
        "trained deep-learning model and the user-facing dashboard. Instead of loading the model "
        "directly inside the dashboard, the dashboard sends requests to FastAPI, and FastAPI returns "
        "the predicted RUL and battery health status."
    )
    doc.add_paragraph(
        "The selected production model served by the API is the Two GLU Blocks + Pooling model. "
        "This model receives a 10-cycle battery window with seven input features and predicts the "
        "remaining useful life in cycles."
    )

    doc.add_heading("2. Purpose of Using FastAPI", level=1)
    doc.add_paragraph(
        "FastAPI was used because it provides a clean, lightweight, and fast way to expose the "
        "trained machine-learning model as a web service. This makes the system more modular: the "
        "model, dashboard, and future external devices can communicate through API endpoints instead "
        "of being tightly coupled in one script."
    )
    add_bullets(
        doc,
        [
            "It separates the machine-learning model from the dashboard interface.",
            "It allows Streamlit or any other frontend to request predictions through HTTP.",
            "It provides automatic API documentation through /docs.",
            "It validates incoming request data using Pydantic models.",
            "It makes the project easier to extend later with XAI, IoT, or model-update endpoints.",
        ],
    )

    doc.add_heading("3. FastAPI Architecture", level=1)
    doc.add_paragraph(
        "The backend is organized into two main files. The first file defines the API endpoints, "
        "while the second file handles model loading, input validation, and prediction logic."
    )
    add_table(
        doc,
        ["File", "Purpose"],
        [
            ["api/main.py", "Defines the FastAPI application, request/response schemas, and API endpoints."],
            ["api/model_service.py", "Loads the GLU model, validates input windows, performs prediction, and classifies the result."],
            ["requirements-api.txt", "Lists the required API packages such as fastapi, uvicorn, tensorflow, numpy, pandas, and pydantic."],
        ],
    )
    doc.add_paragraph("The backend architecture can be summarized as follows:")
    add_code(
        doc,
        "Dashboard or client -> FastAPI endpoint -> Input validation -> GLU model -> Predicted RUL -> Health status -> Response",
    )

    doc.add_heading("4. Model Served by FastAPI", level=1)
    doc.add_paragraph("The API serves the selected GLU model from the following path:")
    add_code(doc, "results/glu_two_blocks_pooling/glu_two_blocks_pooling_model.keras")
    doc.add_paragraph("The model expects a battery window with the following shape:")
    add_code(doc, "(10, 7)")
    add_table(
        doc,
        ["Input Component", "Meaning"],
        [
            ["10", "Sequence length: ten battery cycles in one window."],
            ["7", "Number of model input features per cycle."],
            ["Output", "Predicted Remaining Useful Life in cycles."],
        ],
    )
    doc.add_paragraph("The feature order expected by the API is:")
    add_code(doc, "chI, chV, chT, disI, disV, BCt, SOH")

    doc.add_heading("5. API Endpoints", level=1)
    add_table(
        doc,
        ["Endpoint", "Method", "Description"],
        [
            ["/", "GET", "Returns a simple message confirming that the Battery RUL Prediction API is running."],
            ["/health", "GET", "Checks API status, model availability, model path, input shape, and feature columns."],
            ["/predict", "POST", "Receives a custom scaled battery window with shape (10, 7) and returns predicted RUL and status."],
            ["/sample/{split}/{index}", "GET", "Returns one saved sample window from train, validation, or test data."],
            ["/predict-sample/{split}/{index}", "GET", "Loads a saved sample window, predicts RUL, and returns prediction with true RUL metadata."],
        ],
    )

    doc.add_heading("6. Request and Response Format", level=1)
    doc.add_paragraph("For custom prediction, the user sends a JSON request to /predict.")
    add_code(
        doc,
        '{\n  "window": [\n    [chI, chV, chT, disI, disV, BCt, SOH],\n    "... 10 rows total ..."\n  ]\n}',
    )
    doc.add_paragraph(
        "The response contains the predicted RUL, the health status, the input shape, feature "
        "columns, window size, and model name."
    )
    add_code(
        doc,
        '{\n  "predicted_rul": 167.56,\n  "status": "healthy",\n  "input_shape": [1, 10, 7],\n  "feature_columns": ["chI", "chV", "chT", "disI", "disV", "BCt", "SOH"],\n  "window_size": 10,\n  "model_name": "Two GLU Blocks + Pooling"\n}',
    )

    doc.add_heading("7. Health Status Logic", level=1)
    doc.add_paragraph(
        "The API converts the predicted RUL value into a simple status that can be understood by "
        "the dashboard user."
    )
    add_table(
        doc,
        ["Predicted RUL", "Status", "Meaning"],
        [
            ["Above 50 cycles", "healthy", "The battery still has enough remaining cycles."],
            ["50 cycles or below", "warning", "The battery is approaching end-of-life and should be monitored."],
            ["20 cycles or below", "critical", "The battery is near failure and should be inspected or replaced."],
        ],
    )

    doc.add_heading("8. Dashboard Integration", level=1)
    doc.add_paragraph(
        "The Streamlit dashboard uses FastAPI as the prediction engine. When the user selects or "
        "uploads battery data, the dashboard prepares a 10-cycle window and sends it to FastAPI. "
        "FastAPI validates the input, calls the GLU model, and returns the result. The dashboard "
        "then displays a user-friendly battery health alert."
    )
    add_code(
        doc,
        "Streamlit Dashboard -> POST /predict -> FastAPI -> GLU model -> Predicted RUL + status -> Dashboard alert",
    )

    doc.add_heading("9. How to Run FastAPI", level=1)
    doc.add_paragraph("First, install the API dependencies:")
    add_code(doc, "python -m pip install -r requirements-api.txt")
    doc.add_paragraph("Then start the FastAPI server:")
    add_code(doc, "python -m uvicorn api.main:app --reload --host 127.0.0.1 --port 8000")
    doc.add_paragraph("After starting the server, the API documentation can be opened at:")
    add_code(doc, "http://127.0.0.1:8000/docs")

    doc.add_heading("10. Challenges Faced During FastAPI Development", level=1)
    add_table(
        doc,
        ["Challenge", "What Happened", "Solution"],
        [
            ["Model file missing after branch merge", "After merging the XAI branch, FastAPI failed on startup because the .keras model file was not found.", "The startup logic was updated to check model availability first. The /health endpoint reports model_available as true or false instead of crashing."],
            ["Dashboard needed to work when model artifacts were missing", "When model/result files were not present, the dashboard could not call live prediction.", "The dashboard was updated to fall back to saved XAI prediction results for official B6 mode while live prediction waits for the model file."],
            ["Input shape mismatch risk", "The GLU model requires exactly (10, 7). Sending one row or the wrong feature count would break prediction.", "A validate_window function checks shape and rejects invalid requests with a clear error message."],
            ["Custom GLU layer loading", "The Keras model uses a custom GLUBlock layer, so loading the saved model requires the custom object definition.", "The GLUBlock class was included in model_service.py and passed to load_model using custom_objects."],
            ["Avoid loading the model repeatedly", "Loading a TensorFlow model for every request would be inefficient.", "The model loading function uses lru_cache so the model is loaded once and reused."],
            ["Need for simple client output", "Raw predicted RUL values are not enough for non-technical users.", "The API classifies predictions into healthy, warning, and critical statuses for dashboard alerts."],
            ["TensorFlow startup warnings", "TensorFlow printed oneDNN and placeholder warnings in the terminal.", "These warnings were treated as non-blocking runtime messages. Prediction still worked correctly."],
        ],
    )

    doc.add_heading("11. Testing and Verification", level=1)
    doc.add_paragraph(
        "The API was tested using FastAPI TestClient and sample windows from the processed dataset. "
        "The /health endpoint returned model information correctly, and /predict-sample/test/0 "
        "returned a valid prediction for battery B6."
    )
    add_table(
        doc,
        ["Test", "Expected Result"],
        [
            ["GET /health", "Returns API status, model path, input shape, and feature columns."],
            ["GET /predict-sample/test/0", "Returns predicted RUL, true RUL, metadata, and status for a saved B6 test window."],
            ["POST /predict with wrong shape", "Returns a clear error explaining the expected shape is (10, 7)."],
        ],
    )

    doc.add_heading("12. Future Improvements", level=1)
    add_bullets(
        doc,
        [
            "Add a dedicated /explain endpoint for local XAI explanations for each uploaded window.",
            "Add a /model-update endpoint to start candidate retraining from uploaded datasets.",
            "Add authentication so only admins can use model-update features.",
            "Store predictions and uploaded readings in a database such as SQLite or PostgreSQL.",
            "Connect the API to ESP32 or a Battery Management System for real-time battery alerts.",
            "Add API versioning to separate stable production endpoints from experimental endpoints.",
        ],
    )

    doc.add_heading("13. Conclusion", level=1)
    doc.add_paragraph(
        "The FastAPI backend is an important part of the Battery RUL prediction system because it "
        "turns the trained GLU model into a reusable service. It allows the dashboard, future "
        "hardware devices, and future external systems to request predictions in a clean and "
        "organized way. The main challenges were related to model availability, input validation, "
        "custom model loading, and keeping the dashboard functional even when some generated "
        "artifacts were missing. These issues were handled by adding model availability checks, "
        "strict shape validation, cached model loading, and clear health/status responses."
    )

    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build_document().resolve())
